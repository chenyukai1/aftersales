"""查询 API（契约与 Mock 版一致：POST /api/method/aftersales.after_sales.api.get_vehicle_info 等）。"""
import frappe


def _find_by(doctype, filters):
    names = frappe.get_all(doctype, filters=filters, limit=1, pluck="name")
    return names[0] if names else None


def _result(status, data=None, message=None):
    return {"status": status, "data": data, "message": message}


@frappe.whitelist()
def get_vehicle_info(key=None, chassis_no=None, serial_no=None):
    """入参：车架号或序列号（key 或 chassis_no / serial_no）。
    返回：{status, data:{vehicle, parts_registration, count}, message}"""
    key = (key or chassis_no or serial_no or "").strip()
    if not key:
        return _result("error", message="参数缺失：请提供 车架号 或 序列号")

    name = _find_by("Vehicle Delivery", {"chassis_no": key}) or _find_by(
        "Vehicle Delivery", {"serial_no": key}
    )
    if not name:
        return _result("not_found", message=f'未找到车架号或序列号为 "{key}" 的整车发货记录')

    vehicle = frappe.get_doc("Vehicle Delivery", name).as_dict()
    parts = frappe.get_all(
        "Special Part Registration",
        filters={"chassis_no": vehicle["chassis_no"]},
        fields=["*"],
    )
    return _result("success", data={"vehicle": vehicle, "parts_registration": parts, "count": len(parts)})


@frappe.whitelist()
def get_part_info(key=None, k3_code=None):
    """入参：K3 编码。返回：{status, data:{part}, message}"""
    key = (key or k3_code or "").strip()
    if not key:
        return _result("error", message="参数缺失：请提供 K3 编码")

    name = _find_by("Spare Part", {"k3_code": key})
    if not name:
        return _result("not_found", message=f'未找到 K3 编码为 "{key}" 的配件记录')

    return _result("success", data={"part": frappe.get_doc("Spare Part", name).as_dict()})

# ============================================================
# 报表 Excel 导出
# ============================================================
EXPORT_COLUMNS = [
    ("name", "登记编号"),
    ("feedback_date", "反馈日期"),
    ("chassis_no", "车架号"),
    ("vehicle_model", "车型"),
    ("customer", "客户"),
    ("contact_person", "对接人"),
    ("service_type", "服务类型"),
    ("after_sale_type", "售后类型"),
    ("fault_description", "故障描述"),
    ("handling_action", "索赔处理动作"),
    ("workflow_state", "审批状态"),
    ("customer_status", "客户状态"),
    ("department_status", "部门状态"),
    ("claim_month", "索赔月份"),
    ("claim_week", "索赔周数"),
]


@frappe.whitelist()
def export_service_requests():
    """售后登记列表导出 Excel（按当前用户可见范围，xlsx 返回浏览器下载）。"""
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    import frappe.utils

    fields = [c[0] for c in EXPORT_COLUMNS if frappe.db.has_column("Service Request", c[0])]
    rows = frappe.get_list(
        "Service Request",
        fields=["name"] + fields,
        order_by="creation desc",
        limit_page_length=0,
    )
    wb = Workbook()
    ws = wb.active
    ws.title = "售后登记"
    headers = ["登记编号"] + [c[1] for c in EXPORT_COLUMNS if c[0] in fields]
    ws.append(headers)
    fill = PatternFill("solid", fgColor="EDF2F9")
    bold = Font(bold=True)
    for cell in ws[1]:
        cell.font = bold
        cell.fill = fill
    for r in rows:
        ws.append([r.get("name")] + [r.get(f) for f in fields])
    # 列宽自适应（截断上限 40）
    for col_idx, h in enumerate(headers, 1):
        width = max(
            [len(str(h))]
            + [len(str(ws.cell(row=i, column=col_idx).value or "")) for i in range(2, ws.max_row + 1)]
        )
        from openpyxl.utils import get_column_letter

        ws.column_dimensions[get_column_letter(col_idx)].width = min(width + 2, 40)

    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    frappe.response["filecontent"] = data
    frappe.response["filedata"] = data
    frappe.response["type"] = "binary"
    frappe.response["filename"] = f"售后登记导出_{frappe.utils.today()}.xlsx"
    frappe.response["doctype"] = None
    return
